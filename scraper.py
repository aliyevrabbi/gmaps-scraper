import sys
import asyncio
import urllib.parse
import re
from playwright.async_api import async_playwright
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def clean_text(s):
    if not s:
        return ""
    # Strip non-printable, unicode directional characters, and question marks
    s = re.sub(r'[\u200e\u200f\u202a\u202b\u202c\u200d?]', '', s)
    return s.strip()

def parse_rating(s):
    s = clean_text(s)
    if not s:
        return ""
    try:
        return float(s.replace(',', '.'))
    except ValueError:
        return s

def parse_reviews_count(s):
    s = clean_text(s)
    if not s:
        return ""
    digits = re.sub(r'\D', '', s)
    return int(digits) if digits else ""

def is_instagram_url(href):
    if not href:
        return False
    return 'instagram.com' in href.lower()

def clean_instagram_url(href):
    if not href:
        return ""
    href = href.strip()
    # Strip tracking query params (?igshid=..., ?utm_...) for a clean profile link
    href = re.split(r'[?#]', href)[0]
    return href

async def scrape_google_maps(query, target_count, output_filename, progress_callback=None, headless=True):
    if not output_filename.endswith('.xlsx'):
        output_filename += '.xlsx'
        
    try:
        target_count = int(target_count)
    except ValueError:
        print("[!] Invalid target count. Defaulting to 10.")
        target_count = 10

    cb = (lambda e, d: progress_callback(e, d)) if progress_callback else (lambda e, d: None)
    
    # Low-RAM Chromium arguments for cloud containers
    chromium_args = [
        '--no-sandbox',
        '--disable-dev-shm-usage',
        '--disable-gpu',
        '--disable-software-rasterizer',
        '--no-zygote',
        '--single-process',
        '--disable-setuid-sandbox',
        '--no-first-run',
        '--lang=en-US'
    ]
        
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=headless, args=chromium_args)
        context = await browser.new_context(viewport={'width': 1280, 'height': 800})
        
        search_page = await context.new_page()
        detail_page = await context.new_page()
        
        # Network interception to block unnecessary resources
        async def handle_route(route):
            resource_type = route.request.resource_type
            if resource_type in ['image', 'stylesheet', 'font', 'media', 'manifest', 'other']:
                await route.abort()
            else:
                await route.continue_()
        
        await search_page.route('**/*', handle_route)
        await detail_page.route('**/*', handle_route)

        print(f"[*] Navigating to Google Maps for query: '{query}'")
        encoded_query = urllib.parse.quote_plus(query)
        await search_page.goto(f"https://www.google.com/maps/search/{encoded_query}", wait_until="domcontentloaded", timeout=60000)

        seen_urls = set()
        queue = []
        results = []
        
        feed_selector = 'div[role="feed"]'
        try:
            await search_page.wait_for_selector(feed_selector, timeout=15000)
        except Exception:
            print("[!] Feed container not found. Check if query returns direct results.")
            
        feed_element = search_page.locator(feed_selector)
        reached_end = False
        no_new_links_count = 0

        print(f"[*] Target: {target_count} website-less businesses.")
        
        while len(results) < target_count and not reached_end:
            cards = search_page.locator('div[role="article"], div.Nv2pk')
            count = await cards.count()
            
            new_found_in_scroll = False
            for i in range(count):
                card = cards.nth(i)
                link = card.locator('a').first
                if await link.count() > 0:
                    href = await link.get_attribute('href')
                    if href and href not in seen_urls:
                        seen_urls.add(href)
                        queue.append(href)
                        new_found_in_scroll = True

            while queue and len(results) < target_count:
                url = queue.pop(0)
                print(f"[*] Inspecting URL... (Collected so far: {len(results)}/{target_count})")
                cb("inspect", None)
                try:
                    await detail_page.goto(url, wait_until='domcontentloaded', timeout=60000)
                    
                    # Ensure the panel is fully rendered before checking for the website
                    try:
                        await detail_page.wait_for_selector('h1', timeout=5000)
                    except Exception:
                        pass 
                        
                    await detail_page.wait_for_timeout(1000) # Brief wait for DOM attributes to populate

                    # Only take businesses where the "Website" field on Google Maps
                    # is itself an Instagram link. Anything else (a real website,
                    # or no website field at all) is skipped.
                    authority_href = ""
                    auth_btn = detail_page.locator('a[data-item-id="authority"]').first
                    if await auth_btn.count() > 0:
                        authority_href = await auth_btn.get_attribute('href') or ""

                    if not authority_href.strip():
                        # No website field at all -> nothing to extract, skip.
                        print("      -> Skipped: No website field found.")
                        cb("skip", None)
                        continue

                    if not is_instagram_url(authority_href):
                        # Website field points to a real website -> skip.
                        print("      -> Skipped: Real website detected.")
                        cb("skip", None)
                        continue

                    instagram_link = clean_instagram_url(authority_href)

                    # Extract Business Name
                    business_name = ""
                    name_loc = detail_page.locator('h1').first
                    if await name_loc.count() > 0:
                        business_name = clean_text(await name_loc.inner_text())

                    # Extract Rating
                    rating = ""
                    rating_loc = detail_page.locator('div.F7nice span[aria-hidden="true"]').first
                    if await rating_loc.count() > 0:
                        rating = parse_rating(await rating_loc.inner_text())

                    # Extract Reviews Count
                    reviews_count = ""
                    reviews_loc = detail_page.locator('div.F7nice span[role="img"], div.F7nice span:has-text("(")').last
                    if await reviews_loc.count() > 0:
                        reviews_count = parse_reviews_count(await reviews_loc.inner_text())

                    # Extract Phone Number
                    phone_number = ""
                    phone_btn = detail_page.locator('button[data-item-id^="phone:tel:"]').first
                    if await phone_btn.count() > 0:
                        phone_div = phone_btn.locator('div.IoGyTe').first
                        if await phone_div.count() > 0:
                            phone_number = clean_text(await phone_div.inner_text())
                        else:
                            raw_id = await phone_btn.get_attribute('data-item-id')
                            if raw_id:
                                phone_number = clean_text(raw_id.replace('phone:tel:', ''))

                    # Extract Address (Updated with fallbacks)
                    address = ""
                    address_btn = detail_page.locator('button[data-item-id="address"]').first
                    
                    if await address_btn.count() > 0:
                        # 1. Primary: Extract inner text from div.IoGyTe
                        address_div = address_btn.locator('div.IoGyTe').first
                        if await address_div.count() > 0:
                            address = clean_text(await address_div.inner_text())
                        
                        # 2. Secondary fallback: Extract aria-label attribute
                        if not address:
                            aria_label = await address_btn.get_attribute('aria-label')
                            if aria_label:
                                address = clean_text(aria_label)
                                address = re.sub(r'^(Address|Adres|Ünvan)[\s:]*', '', address, flags=re.IGNORECASE)

                    # 3. Third fallback: Target tooltips if main button fails
                    if not address:
                        tooltip_loc = detail_page.locator('[data-tooltip*="address"], [data-tooltip*="adres"], [data-tooltip*="ünvan"], [data-tooltip*="Copy address"]').first
                        if await tooltip_loc.count() > 0:
                            tooltip_div = tooltip_loc.locator('div.IoGyTe').first
                            if await tooltip_div.count() > 0:
                                address = clean_text(await tooltip_div.inner_text())
                            if not address:
                                aria_label = await tooltip_loc.get_attribute('aria-label')
                                if aria_label:
                                    address = clean_text(aria_label)
                                    address = re.sub(r'^(Address|Adres|Ünvan)[\s:]*', '', address, flags=re.IGNORECASE)

                    if business_name:
                        print(f"      -> SUCCESS: No website, Instagram found. Extracted: {business_name}")
                        
                        result_dict = {
                            "Business Name": business_name,
                            "Rating": rating,
                            "Reviews Count": reviews_count,
                            "Phone Number": phone_number,
                            "Instagram": instagram_link,
                            "Address": address,
                            "URL": url
                        }
                        results.append(result_dict)
                        cb("result", result_dict)

                except Exception as e:
                    print(f"      -> Error inspecting URL: skipped.")
                    cb("inspect_error", str(e))

            if len(results) >= target_count:
                print(f"[*] Target of {target_count} reached.")
                break

            if new_found_in_scroll:
                no_new_links_count = 0
            else:
                no_new_links_count += 1
                
            end_text = search_page.get_by_text("You've reached the end of the list", exact=False)
            if await end_text.count() > 0:
                print("[*] Reached the end of the Google Maps results list.")
                reached_end = True
                break

            if no_new_links_count >= 3:
                print("[*] No new unique links found after 3 consecutive scrolls. Ending search.")
                reached_end = True
                break

            print("[*] Scrolling for more results...")
            cb("scroll", None)
            try:
                await feed_element.evaluate('(node) => node.scrollBy(0, 5000)')
            except Exception:
                pass
            await search_page.wait_for_timeout(2500)

        await browser.close()

        # Excel Export with openpyxl
        wb = Workbook()
        ws = wb.active
        ws.title = "Google Maps Leads"

        headers = ["Business Name", "Rating", "Reviews Count", "Phone Number", "Instagram", "Address"]
        ws.append(headers)

        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        link_font = Font(name="Segoe UI", size=10, color="0563C1", underline="single")
        data_font = Font(name="Segoe UI", size=10, color="1A1A1A")

        for row_data in results:
            row_idx = ws.max_row + 1
            
            ws.cell(row=row_idx, column=1, value=row_data['Business Name'])
            ws.cell(row=row_idx, column=2, value=row_data['Rating'])
            ws.cell(row=row_idx, column=3, value=row_data['Reviews Count'])
            ws.cell(row=row_idx, column=4, value=row_data['Phone Number'])
            ws.cell(row=row_idx, column=5, value=row_data['Instagram'])
            ws.cell(row=row_idx, column=6, value=row_data['Address'])

            name_cell = ws.cell(row=row_idx, column=1)
            name_cell.hyperlink = row_data['URL']
            name_cell.font = link_font
            name_cell.alignment = left_align

            for col_idx in range(2, 7):
                ws.cell(row=row_idx, column=col_idx).font = data_font

            ws.cell(row=row_idx, column=2).alignment = center_align
            ws.cell(row=row_idx, column=3).alignment = center_align
            ws.cell(row=row_idx, column=4).alignment = left_align
            ws.cell(row=row_idx, column=5).alignment = left_align
            ws.cell(row=row_idx, column=6).alignment = left_align

            # Make the Instagram cell itself a clickable hyperlink
            insta_cell = ws.cell(row=row_idx, column=5)
            if row_data['Instagram']:
                insta_cell.hyperlink = row_data['Instagram']
                insta_cell.font = link_font

        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    cell_val = str(cell.value) if cell.value is not None else ""
                    if len(cell_val) > max_length:
                        max_length = len(cell_val)
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = max_length + 3

        ws.views.sheetView[0].showGridLines = True
        
        try:
            wb.save(output_filename)
            print(f"\n[*] FINISHED. Extracted {len(results)} places WITHOUT websites but WITH Instagram pages.")
            print(f"[*] Data successfully saved with hyperlinks and formatting to '{output_filename}'.")
            cb("done", {"total": len(results), "filename": output_filename})
        except Exception as e:
            print(f"\n[!] Error saving Excel file: {e}")
            cb("error", str(e))

if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Usage: python scraper.py <query> <target_count> <output_filename>")
        sys.exit(1)
    
    query_arg = sys.argv[1]
    target_arg = sys.argv[2]
    out_arg = sys.argv[3]
    
    asyncio.run(scrape_google_maps(query_arg, target_arg, out_arg))
