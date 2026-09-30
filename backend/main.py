from __future__ import annotations

import asyncio
import re
import threading
import urllib.parse
import uuid
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from playwright.async_api import async_playwright
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_DIR = Path(__file__).resolve().parent / "outputs"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
executor = ThreadPoolExecutor(max_workers=2)

app = FastAPI(title="LeadZen Google Maps Scraper API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


class CreateJobPayload(BaseModel):
    query: str = Field(min_length=1, max_length=300)
    location: str = Field(default="", max_length=200)
    targetCount: int = Field(default=20, ge=1, le=100)
    filename: str = Field(default="leads.xlsx", max_length=120)


@dataclass
class Job:
    job_id: str
    payload: CreateJobPayload
    status: str = "running"
    inspected: int = 0
    skipped: int = 0
    found: int = 0
    errors: int = 0
    message: str = "Hazırlanır..."
    results: list[dict[str, Any]] = field(default_factory=list)
    output_path: Path | None = None
    cancel_event: threading.Event = field(default_factory=threading.Event)
    future: Any = None

    def status_json(self) -> dict[str, Any]:
        return {
            "jobId": self.job_id,
            "status": self.status,
            "inspected": self.inspected,
            "skipped": self.skipped,
            "found": self.found,
            "errors": self.errors,
            "message": self.message,
        }


jobs: dict[str, Job] = {}
jobs_lock = threading.Lock()


def clean_text(value: str | None) -> str:
    if not value:
        return ""
    return re.sub(r"[\u200e\u200f\u202a\u202b\u202c\u200d?]", "", value).strip()


def parse_rating(value: str | None) -> float | str:
    value = clean_text(value)
    if not value:
        return ""
    try:
        return float(value.replace(",", "."))
    except ValueError:
        return value


def parse_reviews(value: str | None) -> int | str:
    value = clean_text(value)
    digits = re.sub(r"\D", "", value)
    return int(digits) if digits else ""


def is_instagram_url(url: str) -> bool:
    return bool(url and "instagram.com" in url.lower())


def clean_instagram_url(url: str) -> str:
    return re.split(r"[?#]", (url or "").strip())[0]


def safe_filename(name: str) -> str:
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", (name or "leads.xlsx")).strip("._") or "leads"
    if not name.lower().endswith((".xlsx", ".xls", ".csv")):
        name += ".xlsx"
    return name


def update(job: Job, event: str, data: Any = None) -> None:
    with jobs_lock:
        if event == "inspect":
            job.inspected += 1
            job.message = "Biznes yoxlanılır..."
        elif event == "skip":
            job.skipped += 1
        elif event == "result":
            job.results.append(data)
            job.found = len(job.results)
            job.message = f"Yeni lead tapıldı: {data.get('businessName', '')}"
        elif event == "scroll":
            job.message = "Daha çox nəticə üçün Google Maps siyahısı sürüşdürülür..."
        elif event == "error":
            job.errors += 1
            job.message = "Keçici səhv baş verdi, davam edilir..."


async def scrape_job(job: Job) -> None:
    query = job.payload.query.strip()
    location = job.payload.location.strip()
    full_query = f"{query} {location}".strip()
    target = job.payload.targetCount
    seen_urls: set[str] = set()
    queue: list[str] = []
    reached_end = False
    no_new_links = 0

    try:
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True, args=["--lang=en-US", "--disable-dev-shm-usage"])
            context = await browser.new_context(viewport={"width": 1280, "height": 800}, locale="en-US")
            search_page = await context.new_page()
            detail_page = await context.new_page()
            encoded = urllib.parse.quote_plus(full_query)
            with jobs_lock:
                job.message = f"Google Maps-də '{full_query}' axtarılır..."
            await search_page.goto(f"https://www.google.com/maps/search/{encoded}", wait_until="domcontentloaded", timeout=45000)
            try:
                await search_page.wait_for_selector('div[role="feed"]', timeout=15000)
            except Exception:
                pass
            feed = search_page.locator('div[role="feed"]')

            while len(job.results) < target and not reached_end and not job.cancel_event.is_set():
                cards = search_page.locator('div[role="article"], div.Nv2pk')
                count = await cards.count()
                new_found = False
                for i in range(count):
                    if job.cancel_event.is_set():
                        break
                    link = cards.nth(i).locator("a").first
                    if await link.count() > 0:
                        href = await link.get_attribute("href")
                        if href and href not in seen_urls:
                            seen_urls.add(href)
                            queue.append(href)
                            new_found = True

                while queue and len(job.results) < target and not job.cancel_event.is_set():
                    url = queue.pop(0)
                    update(job, "inspect")
                    try:
                        await detail_page.goto(url, wait_until="domcontentloaded", timeout=30000)
                        try:
                            await detail_page.wait_for_selector("h1", timeout=7000)
                        except Exception:
                            pass
                        await detail_page.wait_for_timeout(700)
                        authority = detail_page.locator('a[data-item-id="authority"]').first
                        authority_href = await authority.get_attribute("href") if await authority.count() else ""
                        if not authority_href or not is_instagram_url(authority_href):
                            update(job, "skip")
                            continue

                        name_loc = detail_page.locator("h1").first
                        name = clean_text(await name_loc.inner_text()) if await name_loc.count() else ""
                        rating_loc = detail_page.locator('div.F7nice span[aria-hidden="true"]').first
                        rating = parse_rating(await rating_loc.inner_text()) if await rating_loc.count() else ""
                        reviews_loc = detail_page.locator('div.F7nice span[role="img"], div.F7nice span:has-text("(")').last
                        reviews = parse_reviews(await reviews_loc.inner_text()) if await reviews_loc.count() else ""
                        phone_btn = detail_page.locator('button[data-item-id^="phone:tel:"]').first
                        phone = ""
                        if await phone_btn.count():
                            phone_div = phone_btn.locator("div.IoGyTe").first
                            if await phone_div.count():
                                phone = clean_text(await phone_div.inner_text())
                            else:
                                raw = await phone_btn.get_attribute("data-item-id") or ""
                                phone = clean_text(raw.replace("phone:tel:", ""))
                        address = ""
                        address_btn = detail_page.locator('button[data-item-id="address"]').first
                        if await address_btn.count():
                            address_div = address_btn.locator("div.IoGyTe").first
                            if await address_div.count():
                                address = clean_text(await address_div.inner_text())
                            if not address:
                                label = await address_btn.get_attribute("aria-label") or ""
                                address = re.sub(r"^(Address|Adres|Ünvan)[\s:]*", "", clean_text(label), flags=re.I)
                        result = {
                            "businessName": name,
                            "rating": rating,
                            "reviewsCount": reviews,
                            "phoneNumber": phone,
                            "instagram": clean_instagram_url(authority_href),
                            "address": address,
                            "googleMapsUrl": url,
                            "scrapedAt": datetime.now().isoformat(timespec="seconds"),
                        }
                        if name:
                            update(job, "result", result)
                    except Exception:
                        update(job, "error")

                if len(job.results) >= target or job.cancel_event.is_set():
                    break
                no_new_links = 0 if new_found else no_new_links + 1
                if await search_page.get_by_text("You've reached the end of the list", exact=False).count():
                    reached_end = True
                    break
                if no_new_links >= 3:
                    reached_end = True
                    break
                update(job, "scroll")
                try:
                    await feed.evaluate("node => node.scrollBy(0, 5000)")
                except Exception:
                    pass
                await search_page.wait_for_timeout(1800)
            await browser.close()

        if job.cancel_event.is_set():
            with jobs_lock:
                job.status = "cancelled"
                job.message = "İstifadəçi tərəfindən dayandırıldı. Hazırkı nəticələr saxlanıldı."
        else:
            job.output_path = write_excel(job)
            with jobs_lock:
                job.status = "completed"
                job.message = f"Tamamlandı! {len(job.results)} uyğun biznes tapıldı."
    except Exception as exc:
        with jobs_lock:
            job.status = "error"
            job.message = f"Scraper xətası: {str(exc)[:240]}"


def write_excel(job: Job) -> Path:
    path = OUTPUT_DIR / f"{job.job_id}_{safe_filename(job.payload.filename)}"
    wb = Workbook()
    ws = wb.active
    ws.title = "Google Maps Leads"
    headers = ["Business Name", "Rating", "Reviews Count", "Phone Number", "Instagram", "Address", "Google Maps URL"]
    ws.append(headers)
    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for row in job.results:
        ws.append([row.get("businessName", ""), row.get("rating", ""), row.get("reviewsCount", ""), row.get("phoneNumber", ""), row.get("instagram", ""), row.get("address", ""), row.get("googleMapsUrl", "")])
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = min(max(len(str(c.value or "")) for c in col) + 3, 60)
    wb.save(path)
    return path


def launch_job(job: Job) -> None:
    asyncio.run(scrape_job(job))


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/jobs")
def create_job(payload: CreateJobPayload) -> dict[str, str]:
    job_id = f"job_{uuid.uuid4().hex[:10]}"
    job = Job(job_id=job_id, payload=payload)
    with jobs_lock:
        jobs[job_id] = job
    job.future = executor.submit(launch_job, job)
    return {"jobId": job_id, "status": "running", "message": "Scraper başladıldı"}


@app.get("/api/jobs/{job_id}")
def get_status(job_id: str) -> dict[str, Any]:
    job = jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="İş tapılmadı")
    with jobs_lock:
        return job.status_json()


@app.get("/api/jobs/{job_id}/results")
def get_results(job_id: str) -> list[dict[str, Any]]:
    job = jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="İş tapılmadı")
    with jobs_lock:
        return list(job.results)


@app.post("/api/jobs/{job_id}/cancel")
def cancel_job(job_id: str) -> dict[str, Any]:
    job = jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="İş tapılmadı")
    job.cancel_event.set()
    return {"success": True, "message": "Scraper dayandırılması tələb olundu"}


@app.get("/api/jobs/{job_id}/export", response_model=None)
def export_job(job_id: str) -> FileResponse:
    job = jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="İş tapılmadı")
    if not job.output_path or not job.output_path.exists():
        job.output_path = write_excel(job)
    return FileResponse(job.output_path, filename=safe_filename(job.payload.filename), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# In the production Docker image the Vite build is copied here so one Render
# Web Service can serve both the mobile UI and the REST API from one URL.
frontend_dir = Path(__file__).resolve().parent.parent / "frontend_dist"
if frontend_dir.exists():
    app.mount("/", StaticFiles(directory=frontend_dir, html=True), name="frontend")
